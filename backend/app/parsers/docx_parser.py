import io
import logging
import docx

logger = logging.getLogger(__name__)


def parse_docx(file_bytes: bytes) -> str:
    """
    Parses DOCX bytes using python-docx. Extracts paragraphs and tables.
    """
    if not file_bytes:
        raise ValueError("DOCX file buffer is empty")

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = []

        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        extracted = "\n".join(paragraphs).strip()

        if not extracted:
            raise ValueError("DOCX document contains no readable text content")

        return extracted
    except Exception as e:
        logger.error(f"DOCX parsing failed: {e}")
        raise ValueError(f"Failed to parse DOCX file: {str(e)}")
